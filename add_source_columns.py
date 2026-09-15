#!/usr/bin/env python3
"""
Creates run_report_v5.0.xlsx from the CURRENT state of run_report_v4.0.xlsx
(including any manual edits already made to it) by appending two new
columns -- 'source' (source name) and 'source_url' -- to the 'All Drivers'
and 'Failed & HITL' sheets.

Deliberately does NOT regenerate the workbook from the underlying registry/
reports the way combine_final_report.py does: v4.0.xlsx on disk is treated
as the authoritative baseline, exactly as it currently is, and only the two
new columns are appended on top of it. Every other sheet is left untouched.

Example
-------
    python add_source_columns.py
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import openpyxl   # noqa: E402

from combine_final_report import (      # noqa: E402
    build_best_known_outcomes, find_fresh_reports, load_report_all_drivers,
)

ROOT = Path(__file__).parent


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--baseline-in", default="output/monthly/run_report_v4.0.xlsx",
                        help="the existing workbook to augment, used as-is (edits and all)")
    parser.add_argument("--baseline-report", default="output/monthly/run_report_v3.0.xlsx")
    parser.add_argument("--curated-dir", default="output/monthly")
    parser.add_argument("--output", default="output/monthly/run_report_v5.0.xlsx")
    parser.add_argument("--sheets", nargs="+", default=["All Drivers", "Failed & HITL"])
    args = parser.parse_args()

    src_path = ROOT / args.baseline_in
    out_path = ROOT / args.output
    curated_dir = ROOT / args.curated_dir

    if not src_path.exists():
        print(f"ERROR: {src_path} does not exist")
        return 1

    shutil.copy(src_path, out_path)
    print(f"Copied {src_path.name} -> {out_path.name} (starting point, edits preserved)")

    # ---- build the same driver_id -> source/source_url lookup combine_final_report.py uses
    baseline_df = load_report_all_drivers(ROOT / args.baseline_report)
    fresh_reports = find_fresh_reports(curated_dir)
    best = build_best_known_outcomes(baseline_df, fresh_reports)
    print(f"Built source/source_url lookup for {len(best)} driver_id(s) "
          f"from run_report_v3.0.xlsx + {len(fresh_reports)} fresh report(s)")

    wb = openpyxl.load_workbook(out_path)
    for sheet_name in args.sheets:
        if sheet_name not in wb.sheetnames:
            print(f"  (skipping '{sheet_name}': not found in {out_path.name})")
            continue
        ws = wb[sheet_name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if "driver_id" not in header:
            print(f"  (skipping '{sheet_name}': no driver_id column)")
            continue
        did_col = header.index("driver_id") + 1

        # Don't duplicate the columns if this script is re-run on an already-augmented file.
        source_col = header.index("source") + 1 if "source" in header else ws.max_column + 1
        url_col = header.index("source_url") + 1 if "source_url" in header else ws.max_column + (
            2 if source_col == ws.max_column + 1 else 1)

        if "source" not in header:
            ws.cell(row=1, column=source_col, value="source")
        if "source_url" not in header:
            ws.cell(row=1, column=url_col, value="source_url")

        filled = 0
        for row_idx in range(2, ws.max_row + 1):
            did = ws.cell(row=row_idx, column=did_col).value
            b = best.get(did)
            if b is None:
                continue
            ws.cell(row=row_idx, column=source_col, value=b.get("source", ""))
            ws.cell(row=row_idx, column=url_col, value=b.get("source_url", ""))
            filled += 1
        print(f"  '{sheet_name}': filled source/source_url for {filled} of {ws.max_row - 1} row(s)")

    wb.save(out_path)
    print(f"\nWrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
