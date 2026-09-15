#!/usr/bin/env python3
"""
Creates run_report_v6.0.xlsx from the current run_report_v5.0.xlsx:

1. In 'All Drivers', column D ('driver') gets its font colored green
   (outcome == 'success') or red (anything else).
2. In 'By Commodity', starting at column G, each commodity's row gets one
   cell per driver belonging to that commodity (from 'All Drivers'),
   sorted alphabetically by driver name, same font-color coding carried
   over -- a transposed, color-coded driver list per commodity.

Example
-------
    python add_commodity_driver_colors.py
"""
from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import openpyxl                     # noqa: E402
from openpyxl.styles import Font    # noqa: E402

ROOT = Path(__file__).parent

GREEN = "FF008000"
RED = "FFFF0000"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--baseline-in", default="output/monthly/run_report_v5.0.xlsx")
    parser.add_argument("--output", default="output/monthly/run_report_v6.0.xlsx")
    args = parser.parse_args()

    src_path = ROOT / args.baseline_in
    out_path = ROOT / args.output

    if not src_path.exists():
        print(f"ERROR: {src_path} does not exist")
        return 1

    shutil.copy(src_path, out_path)
    print(f"Copied {src_path.name} -> {out_path.name}")

    wb = openpyxl.load_workbook(out_path)

    # ---- 1. color column D ('driver') in All Drivers by outcome ----
    ws_all = wb["All Drivers"]
    header = [c.value for c in next(ws_all.iter_rows(min_row=1, max_row=1))]
    idx = {h: i + 1 for i, h in enumerate(header)}  # 1-based column numbers
    commodity_col = idx["commodity"]
    driver_col = idx["driver"]
    outcome_col = idx["outcome"]

    by_commodity: dict[str, list[tuple[str, bool]]] = {}
    colored = 0
    for row_idx in range(2, ws_all.max_row + 1):
        driver_name = ws_all.cell(row=row_idx, column=driver_col).value
        outcome = ws_all.cell(row=row_idx, column=outcome_col).value
        commodity = ws_all.cell(row=row_idx, column=commodity_col).value
        if driver_name is None:
            continue
        is_success = outcome == "success"
        cell = ws_all.cell(row=row_idx, column=driver_col)
        cell.font = Font(color=GREEN if is_success else RED)
        colored += 1
        by_commodity.setdefault(commodity, []).append((driver_name, is_success))

    print(f"Colored {colored} driver-name cell(s) in 'All Drivers' column D "
          f"({sum(1 for v in by_commodity.values() for n, s in v if s)} green, "
          f"{sum(1 for v in by_commodity.values() for n, s in v if not s)} red)")

    # ---- 2. paste each commodity's driver list, transposed, from column G ----
    ws_bc = wb["By Commodity"]
    bc_header = [c.value for c in next(ws_bc.iter_rows(min_row=1, max_row=1))]
    bc_commodity_col = bc_header.index("commodity") + 1

    max_drivers = 0
    for row_idx in range(2, ws_bc.max_row + 1):
        commodity = ws_bc.cell(row=row_idx, column=bc_commodity_col).value
        drivers = sorted(by_commodity.get(commodity, []), key=lambda t: t[0].lower())
        max_drivers = max(max_drivers, len(drivers))
        for offset, (driver_name, is_success) in enumerate(drivers):
            col = 7 + offset  # column G = 7
            cell = ws_bc.cell(row=row_idx, column=col, value=driver_name)
            cell.font = Font(color=GREEN if is_success else RED)

    print(f"Pasted per-commodity driver lists into 'By Commodity' starting at column G "
          f"(widest commodity has {max_drivers} drivers -> column {openpyxl.utils.get_column_letter(6 + max_drivers)})")

    wb.save(out_path)
    print(f"\nWrote {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
